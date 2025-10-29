# =====================================
# Streamlit App: 提成表多sheet自动审核（总 + 轻卡 + 重卡）
# 标红错误格 + 标黄合同号 + 精简错误下载 + 独立错误数统计
# =====================================
import streamlit as st
import pandas as pd
from io import BytesIO
import unicodedata, re

try:
    from openpyxl import Workbook
    from openpyxl.styles import PatternFill
except ImportError:
    st.error("❌ openpyxl 未安装，请执行 pip install openpyxl")
    st.stop()

st.title("📊 提成表多sheet自动审核工具（总 + 轻卡 + 重卡）")

# ========== 上传文件 ==========
uploaded_files = st.file_uploader(
    "请上传包含“提成”、“放款明细”、“二次明细”和“原表”的xlsx文件",
    type="xlsx", accept_multiple_files=True
)

if not uploaded_files or len(uploaded_files) < 4:
    st.warning("⚠️ 请至少上传 提成表、放款明细、二次明细、原表 四个文件")
    st.stop()

# ========== 工具函数 ==========
def find_file(files_list, keyword):
    for f in files_list:
        if keyword in f.name:
            return f
    return None

def normalize_text(val):
    if pd.isna(val):
        return ""
    s = str(val)
    s = re.sub(r'[\n\r\t ]+', '', s)
    s = s.replace('\u3000', '')
    return ''.join(unicodedata.normalize('NFKC', ch) for ch in s).lower().strip()

def normalize_num(val):
    if pd.isna(val):
        return None
    s = str(val).replace(",", "").replace("%", "").strip()
    if s in ["", "-", "nan"]:
        return None
    try:
        return float(s)
    except:
        return None

def find_col(df_like, keyword, exact=False):
    key = keyword.strip().lower()
    columns = df_like.columns if hasattr(df_like, "columns") else df_like.index
    for col in columns:
        cname = str(col).strip().lower()
        if (exact and cname == key) or (not exact and key in cname):
            return col
    return None

# ========== 读取文件 ==========
tc_file = find_file(uploaded_files, "提成")
fk_file = find_file(uploaded_files, "放款明细")
ec_file = find_file(uploaded_files, "二次明细")
original_file = find_file(uploaded_files, "原表")

tc_xls = pd.ExcelFile(tc_file)
sheet_total = next((s for s in tc_xls.sheet_names if "总" in s), None)
sheets_qk = [s for s in tc_xls.sheet_names if "轻卡" in s]
sheets_zk = [s for s in tc_xls.sheet_names if "重卡" in s]

tc_sheets = {
    "总": [pd.read_excel(tc_file, sheet_name=sheet_total)] if sheet_total else [],
    "轻卡": [pd.read_excel(tc_file, sheet_name=s) for s in sheets_qk],
    "重卡": [pd.read_excel(tc_file, sheet_name=s) for s in sheets_zk],
}

fk_xls = pd.ExcelFile(fk_file)
fk_dfs = [pd.read_excel(fk_file, sheet_name=s) for s in fk_xls.sheet_names if "潮掣" in s]

ec_xls = pd.ExcelFile(ec_file)
ec_df = pd.concat([pd.read_excel(ec_file, sheet_name=s) for s in ec_xls.sheet_names], ignore_index=True)

original_df = pd.read_excel(original_file)

st.success(f"✅ 提成表已读取：总({len(tc_sheets['总'])})、轻卡({len(tc_sheets['轻卡'])})、重卡({len(tc_sheets['重卡'])})")

# ========== 定义映射 ==========
MAPPING = {
    "放款日期": ("放款明细", "放款日期", 0, 1),
    "提报人员": ("放款明细", "提报人员", 0, 1),
    "城市经理": ("放款明细", "城市经理", 0, 1),
    "租赁本金": ("放款明细", "租赁本金", 0, 1),
    "收益率": ("放款明细", "xirr", 0.005, 1),
    "期限": ("放款明细", "租赁期限/年", 0.5, 12),
    "家访": ("放款明细", "家访", 0, 1),
    "人员类型": ("放款明细", "类型", 0, 1),
    "二次交接": ("二次明细", "出本流程时间", 0, 1),
}

# ========== 比对函数 ==========
def get_ref_row(contract_no, source_type):
    contract_no = str(contract_no).strip()
    if source_type == "放款明细":
        for df in fk_dfs:
            col = find_col(df, "合同")
            if col is None:
                continue
            res = df[df[col].astype(str).str.strip() == contract_no]
            if not res.empty:
                return res.iloc[0]
    elif source_type == "二次明细":
        col = find_col(ec_df, "合同")
        if col is not None:
            res = ec_df[ec_df[col].astype(str).str.strip() == contract_no]
            if not res.empty:
                return res.iloc[0]
    elif source_type == "原表":
        col = find_col(original_df, "合同", exact=False)
        if col is not None:
            res = original_df[original_df[col].astype(str).str.strip() == contract_no]
            if not res.empty:
                return res.iloc[0]
    return None


# ========== 核心审核函数 ==========
def audit_one_sheet(tc_df, sheet_label):
    contract_col_main = find_col(tc_df, "合同")
    if not contract_col_main:
        st.warning(f"⚠️ {sheet_label}：未找到‘合同’列，跳过。")
        return None, None, 0, 0

    wb = Workbook()
    ws = wb.active
    for i, col_name in enumerate(tc_df.columns, start=1):
        ws.cell(1, i, col_name)

    red_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
    yellow_fill = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")

    total_errors = 0
    error_rows = set()
    n = len(tc_df)

    progress = st.progress(0)
    status = st.empty()
    person_type_col = find_col(tc_df, "人员类型", exact=True)

    for idx, row in tc_df.iterrows():
        contract_no = row.get(contract_col_main)
        if pd.isna(contract_no):
            continue

        row_has_error = False

        for main_kw, (src, ref_kw, tol, mult) in MAPPING.items():
            exact_main = "期限" in main_kw or main_kw == "人员类型"
            main_col = find_col(tc_df, main_kw, exact=exact_main)
            if not main_col:
                continue

            if main_kw == "收益率":
                person_type = str(row.get(person_type_col, "")).strip()
                if person_type == "轻卡":
                    ref_row = get_ref_row(contract_no, "原表")
                    ref_kw = "年化nim"
                else:
                    ref_row = get_ref_row(contract_no, src)
            else:
                ref_row = get_ref_row(contract_no, src)

            if ref_row is None:
                continue

            ref_col = find_col(ref_row, ref_kw, exact=(main_kw == "人员类型"))
            if not ref_col:
                continue

            main_val = row[main_col]
            ref_val = ref_row[ref_col]

            if "日期" in main_kw or main_kw == "二次交接":
                try:
                    main_dt = pd.to_datetime(main_val, errors='coerce').normalize()
                    ref_dt = pd.to_datetime(ref_val, errors='coerce').normalize()
                except:
                    main_dt = ref_dt = pd.NaT
                if pd.isna(main_dt) or pd.isna(ref_dt) or main_dt != ref_dt:
                    row_has_error = True
                    total_errors += 1
                    ws.cell(idx + 2, list(tc_df.columns).index(main_col) + 1).fill = red_fill
            else:
                m = normalize_num(main_val)
                r = normalize_num(ref_val)
                if main_kw == "收益率" and m is not None and r is not None:
                    if m > 1: m /= 100
                    if r > 1: r /= 100
                if m is not None and r is not None:
                    if "期限" in main_kw:
                        r *= mult
                    if abs(m - r) > tol:
                        row_has_error = True
                        total_errors += 1
                        ws.cell(idx + 2, list(tc_df.columns).index(main_col) + 1).fill = red_fill
                else:
                    if normalize_text(main_val) != normalize_text(ref_val):
                        row_has_error = True
                        total_errors += 1
                        ws.cell(idx + 2, list(tc_df.columns).index(main_col) + 1).fill = red_fill

        if row_has_error:
            ws.cell(idx + 2, list(tc_df.columns).index(contract_col_main) + 1).fill = yellow_fill
            error_rows.add(idx)

        for j, val in enumerate(row, start=1):
            ws.cell(idx + 2, j, val)

        if (idx + 1) % 10 == 0 or (idx + 1) == n:
            progress.progress((idx + 1) / n)
            status.text(f"{sheet_label} 审核进度：{idx + 1}/{n}")

    # ===== 输出区 =====
    output_full = BytesIO()
    wb.save(output_full)
    output_full.seek(0)

    # 精简错误表
    wb_err = Workbook()
    ws_err = wb_err.active
    for i, col_name in enumerate(tc_df.columns, start=1):
        ws_err.cell(1, i, col_name)
    row_idx = 2
    for idx in sorted(error_rows):
        for j, val in enumerate(tc_df.iloc[idx], start=1):
            ws_err.cell(row_idx, j, val)
            orig_cell = ws.cell(idx + 2, j)
            fill = orig_cell.fill
            if hasattr(fill, "fill_type") and fill.fill_type == "solid":
                new_fill = PatternFill(fill_type="solid",
                                       start_color=fill.start_color.rgb,
                                       end_color=fill.end_color.rgb)
                ws_err.cell(row_idx, j).fill = new_fill
        row_idx += 1

    output_err = BytesIO()
    wb_err.save(output_err)
    output_err.seek(0)

    return output_full, output_err, total_errors, len(error_rows)

# ========== 审核所有 sheet ==========
results = {}
for label, df_list in tc_sheets.items():
    if not df_list:
        continue
    for i, df in enumerate(df_list, start=1):
        tag = f"{label}{i if len(df_list) > 1 else ''}"
        st.divider()
        st.subheader(f"📘 正在审核：{tag}")
        full, err, errs, rows = audit_one_sheet(df, tag)
        results[tag] = (full, err, errs, rows)

# ========== 🔍 反向漏填检查（保持不变） ==========
st.divider()
st.subheader("🔍 反向漏填检查（仅基于放款明细中包含“潮掣”的sheet）")

def get_contracts_from_df(df):
    col = find_col(df, "合同", exact=False)
    if col is not None:
        return set(df[col].dropna().astype(str).str.strip())
    return set()

def get_contracts_from_fk_dfs(fk_list):
    all_cons = set()
    for df in fk_list:
        all_cons |= get_contracts_from_df(df)
    return all_cons

contracts_total = get_contracts_from_df(tc_sheets["总"][0])
contracts_fk = get_contracts_from_fk_dfs(fk_dfs)
missing_contracts = sorted(list(contracts_fk - contracts_total))

if missing_contracts:
    st.warning(f"⚠️ 发现 {len(missing_contracts)} 个合同号存在于放款明细中，但未出现在提成表‘总’sheet中")
    df_missing = pd.DataFrame({"漏填合同号": missing_contracts})
    output_missing = BytesIO()
    with pd.ExcelWriter(output_missing, engine="openpyxl") as writer:
        df_missing.to_excel(writer, index=False, sheet_name="漏填合同号")
    output_missing.seek(0)
    st.download_button(
        "📥 下载漏填合同号表（基于放款明细-潮掣）",
        data=output_missing,
        file_name="提成_漏填合同号_基于放款明细_潮掣.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
else:
    st.success("✅ 未发现漏填合同号（基于放款明细-潮掣）。")

# ========== 下载区 ==========
st.divider()
st.subheader("📤 下载审核结果文件")

for tag, (full, err, errs, rows) in results.items():
    st.write(f"📘 **{tag}**：发现 {errs} 个错误，共 {rows} 行异常")
    st.download_button(
        f"📥 下载 {tag} 审核标注版",
        data=full,
        file_name=f"提成_{tag}_审核标注版.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    if rows > 0:
        st.download_button(
            f"📥 下载 {tag} 错误精简版（含红黄标记）",
            data=err,
            file_name=f"提成_{tag}_错误精简版.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )

st.success("✅ 所有sheet审核完成！")
